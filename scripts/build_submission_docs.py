"""Build the three submission documents for TECHATHON 2026."""
from __future__ import annotations

from pathlib import Path

from docx import Document
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

ROOT = Path(__file__).resolve().parent.parent
OUT = ROOT / "提交材料"
OUT.mkdir(exist_ok=True)

INK = RGBColor(0x11, 0x11, 0x11)
MUTED = RGBColor(0x66, 0x66, 0x66)
RED = RGBColor(0xE4, 0x00, 0x2B)
FONT = "微软雅黑"
HEAD_FILL = "F2F2F2"


def setup(doc: Document) -> None:
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin = Cm(2.0)
    sec.bottom_margin = Cm(1.8)
    normal = doc.styles["Normal"]
    normal.font.name = "Calibri"
    normal.font.size = Pt(10.5)
    normal._element.rPr.rFonts.set(qn("w:eastAsia"), FONT)


def run_font(run, size=10.5, bold=False, color=INK, font=FONT):
    run.font.name = "Calibri"
    run._element.get_or_add_rPr().rFonts.set(qn("w:eastAsia"), font)
    run.font.size = Pt(size)
    run.font.bold = bold
    run.font.color.rgb = color


def para(doc, text, size=10.5, bold=False, color=INK, after=4, indent=0,
         align=WD_ALIGN_PARAGRAPH.JUSTIFY):
    p = doc.add_paragraph()
    p.alignment = align
    p.paragraph_format.space_after = Pt(after)
    p.paragraph_format.line_spacing = 1.3
    if indent:
        p.paragraph_format.first_line_indent = Cm(indent)
    r = p.add_run(text)
    run_font(r, size, bold, color)
    return p


def bullet(doc, text, size=10.5):
    p = doc.add_paragraph(style="List Bullet")
    p.paragraph_format.space_after = Pt(2)
    p.paragraph_format.line_spacing = 1.25
    r = p.add_run(text)
    run_font(r, size)
    return p


def h1(doc, text, num=None):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(16)
    p.paragraph_format.space_after = Pt(6)
    p.paragraph_format.keep_with_next = True
    if num is not None:
        r = p.add_run(f"{num}  ")
        run_font(r, 17, True, RED)
    r = p.add_run(text)
    run_font(r, 15, True, INK)
    pPr = p._p.get_or_add_pPr()
    pBdr = OxmlElement("w:pBdr")
    bottom = OxmlElement("w:bottom")
    bottom.set(qn("w:val"), "single")
    bottom.set(qn("w:sz"), "8")
    bottom.set(qn("w:space"), "4")
    bottom.set(qn("w:color"), "E4002B")
    pBdr.append(bottom)
    pPr.append(pBdr)
    return p


def h2(doc, text):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(10)
    p.paragraph_format.space_after = Pt(3)
    p.paragraph_format.keep_with_next = True
    r = p.add_run(text)
    run_font(r, 11.5, True, INK)
    return p


def shade(cell, fill):
    tcPr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tcPr.append(shd)


def table(doc, headers, rows, widths, font_size=9):
    t = doc.add_table(rows=1, cols=len(headers))
    t.style = "Table Grid"
    t.autofit = False
    for i, htext in enumerate(headers):
        c = t.rows[0].cells[i]
        c.width = widths[i]
        c.text = ""
        p = c.paragraphs[0]
        r = p.add_run(htext)
        run_font(r, font_size, True)
        shade(c, HEAD_FILL)
    for row in rows:
        cells = t.add_row().cells
        for i, val in enumerate(row):
            cells[i].width = widths[i]
            cells[i].text = ""
            p = cells[i].paragraphs[0]
            p.paragraph_format.space_after = Pt(1)
            r = p.add_run(str(val))
            run_font(r, font_size)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return t


def header_footer(doc, label):
    sec = doc.sections[0]
    hp = sec.header.paragraphs[0]
    hp.text = ""
    r = hp.add_run("MOLVISION  ·  TECHATHON 2026")
    run_font(r, 8, True, MUTED)
    hp.add_run("\t")
    r = hp.add_run(label)
    run_font(r, 8, False, MUTED)
    hp.paragraph_format.tab_stops.add_tab_stop(Cm(16.6), 2)  # right align

    fp = sec.footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    r = fp.add_run("ShanghaiTech TECHATHON 2026  ·  MolVision")
    run_font(r, 7.5, False, MUTED)


def doc_title(doc, title, subtitle, tag):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(40)
    r = p.add_run(tag)
    run_font(r, 9, True, RED)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(8)
    r = p.add_run(title)
    run_font(r, 28, True, INK)
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    run_font(r, 13, False, MUTED)
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(20)
    r = p.add_run("团队：MolVision　|　版本 v1.0　|　2026-10-18")
    run_font(r, 9.5, False, MUTED)


# ---------------------------------------------------------------------------
# 01 说明文档
# ---------------------------------------------------------------------------
def build_intro() -> None:
    doc = Document()
    setup(doc)
    header_footer(doc, "说明文档")
    doc_title(doc, "MolVision 结构式识图", "从一张结构图到可计算的分子", "TECHATHON 2026  ·  说明文档")

    h1(doc, "项目简介", "01")
    para(doc,
         "MolVision 是一个本地运行的化学结构式识图工作台。输入论文截图、专利图、手绘或整页 PDF，"
         "系统输出标准 SMILES、分子描述符、2D 重绘图与 3D 构象，并支持多结构切分、Markush/R 基、"
         "常见缩写基团和 SDF/CSV 导出。它解决的不是“识别一张图”，而是把图片里的结构真正变成"
         "可搜索、可计算、可复现的数据。")
    para(doc,
         "与端到端黑箱不同，MolVision 把识别模型放进一条可检查的化学流水线：切分 → 识别 → "
         "化学归一 → RDKit 校验 → 2D/3D/描述符 → 导出。每一步都有失败兜底和人工介入点，"
         "识别模型可以替换，下游校验和导出的行为保持稳定。")

    h1(doc, "成员信息", "02")
    para(doc, "下表为提交模板，请按实际团队情况补全（当前仅填写了已知成员）。", size=9, color=MUTED)
    table(doc,
          ["姓名", "学校 / 专业", "年级", "分工", "联系方式"],
          [
              ["何宇辉", "【待填写】", "【待填写】", "识别流水线、化学归一、评测", "【待填写】"],
              ["【成员 2】", "【待填写】", "【待填写】", "前端 / Demo / 材料", "【待填写】"],
              ["【成员 3】", "【待填写】", "【待填写】", "【待填写】", "【待填写】"],
          ],
          [Cm(2.6), Cm(4.2), Cm(2.0), Cm(5.0), Cm(2.8)])
    para(doc, "团队名称：MolVision（如主办方要求实名，请替换为注册队名）。",
         size=9, color=MUTED)

    h1(doc, "功能说明", "03")
    table(doc,
          ["功能", "说明"],
          [
              ["单图识别", "PNG / JPG / 手绘截图 → SMILES，失败自动切换手绘模型与旋转重试"],
              ["多结构切分", "一张图里的多个结构自动切分、逐块识别、按阅读顺序输出"],
              ["PDF 批量", "整页 PDF 逐页判断“是否图像页”，输出 CSV / SDF / 汇总 JSON"],
              ["Markush / R 基", "R、R1、R''、Ar 统一映射为 *；未定义时按甲基（CH3）代入生成 3D 与描述符"],
              ["缩写基团", "Ph、Me、Et、iPr、tBu、Bn、Boc、Cbz、Ts、TBS 等 30+ 缩写自动展开"],
              ["化学校验", "RDKit canonical SMILES、InChIKey、价键/芳香性检查、结构警示"],
              ["2D / 3D", "2D 重绘；ETKDGv3 + MMFF94 生成 3D，可选 1/5/10/20 构象"],
              ["描述符", "分子式、MW、logP、TPSA、HBD/HBA、可旋转键、SAscore 等"],
              ["导出", "SDF（含多构象与能量字段）、CSV、批量去重"],
              ["人工介入", "Crop 手动框选、手动输入 SMILES、逐条核对警示"],
          ],
          [Cm(3.2), Cm(12.8)])

    h1(doc, "原创声明", "04")
    para(doc,
         "1. 本项目提交的代码（app/、web/、scripts/ 及配套文档）由团队独立完成，未复制、"
         "未改写他人项目代码；引用或调用的第三方组件已在《第三方资源说明》中逐项列出。")
    para(doc,
         "2. 识别模型使用 DECIMER 2.8 预训练权重，通过其官方 Python 包调用，团队未修改模型结构；"
         "3D 构象、描述符、校验与导出由 RDKit 等开源库完成。")
    para(doc,
         "3. 评测所用分子结构来自公开数据集（RDKit 自带 NCI 子集、PDB 公开条目），测试图片为"
         "团队自行渲染或自绘；未使用未授权的他人图片作为提交资产。")
    para(doc,
         "4. 项目代码建议以 MIT 许可证开源；第三方组件遵循各自许可证，详见《第三方资源说明》。")
    para(doc, "成员签名：____________________　　日期：2026 年 10 月 18 日", size=10)

    doc.save(OUT / "01-说明文档.docx")
    print("saved 01")


# ---------------------------------------------------------------------------
# 02 技术文档
# ---------------------------------------------------------------------------
def build_tech() -> None:
    doc = Document()
    setup(doc)
    header_footer(doc, "技术文档")
    doc_title(doc, "技术文档", "架构、关键算法与评测方法", "TECHATHON 2026  ·  技术文档")

    h1(doc, "系统概述", "01")
    para(doc,
         "MolVision 采用“识别模型 + 化学流水线”的分层架构。识别模型只负责把像素读成 SMILES，"
         "切分、归一、校验、3D 与导出由确定性代码完成。这样做的代价是更多工程代码，收益是"
         "失败可定位、模型可替换、结果可复现。")
    para(doc, "流水线：图片 → 预处理 → 投影切分 → DECIMER 识别 → 化学归一 → RDKit 校验 → "
              "2D 重绘 / 3D 构象 / 描述符 → SDF / CSV。")
    table(doc,
          ["模块", "职责", "实现"],
          [
              ["图像预处理", "EXIF 校正、超大图缩放、二值化", "Pillow + OpenCV"],
              ["结构切分", "投影切分、表格线/箭头过滤、标签重挂", "app/segment.py"],
              ["结构识别", "DECIMER 2.8；失败回退手绘模型；90°/270° 旋转", "app/recognize.py"],
              ["化学归一", "R 基、缩写、同位素氢、盐/溶剂处理", "app/chem.py"],
              ["校验与描述符", "canonical SMILES、InChIKey、警示、SAscore", "RDKit 2026.09.1"],
              ["3D 构象", "ETKDGv3 嵌入 + MMFF94/UFF，多构象", "RDKit"],
              ["PDF 批量", "页面筛选、逐页识别、去重导出", "app/pdf.py + PyMuPDF"],
              ["交付层", "Web 单页、REST API、SDF/CSV", "app/server.py + web/index.html"],
          ],
          [Cm(3.2), Cm(8.8), Cm(4.6)])

    h1(doc, "关键算法", "02")
    h2(doc, "2.1 结构切分")
    para(doc, "对二值化图像分别做行、列方向投影，按空白带切出候选块；用长笔画比例过滤表格线、"
              "箭头和文字；对标签被切断的情况按长笔画外接框与各向异性边距重新挂回。"
              "密集图或切分失败时，UI 提供 Crop 手动框选。")
    h2(doc, "2.2 识别与兜底")
    para(doc, "每个裁剪块送入 DECIMER 2.8。若解析失败，自动切换手绘模型；若仍失败，尝试 90°/270° "
              "旋转。识别结果先做 Markush/缩写/D→H 归一，再过 RDKit 解析与噪声规则（片段数、"
              "重复片段、长烷烃等）。中央 S/P 标签把分子切成两半时，检测两半相邻且含 I/Br 幻觉，"
              "合并后重新识别。")
    h2(doc, "2.3 化学归一规则")
    bullet(doc, "R/R1/R''/Ar → *；未定义 R 在 3D 与描述符阶段按 CH3 代入，界面标注 R assumed: CH3。")
    bullet(doc, "30+ 常见缩写展开；环内展开使用独立 %90+ 环编号，避免与宿主环编号冲突。")
    bullet(doc, "D/²H 在识别路径折叠为 H；手动输入的 [2H] 保留。")
    bullet(doc, "多片段分子保留最大片段（去盐/溶剂），并在描述符中标注 fragments。")
    h2(doc, "2.4 3D 构象")
    para(doc, "ETKDGv3 采样，MMFF94 优化（金属回退 UFF）；单构象模式嵌入 1 个，多构象模式按 "
              "能量排序、保留最低能 +10 kcal/mol 窗口、以 0.5 Å RMSD 去重。SDF 每个构象写入 "
              "ENERGY_KCAL 字段。")
    h2(doc, "2.5 并发与缓存")
    para(doc, "RDKit 分析（InChI、绘图、嵌入）共享全局状态，使用全局锁串行化；DECIMER 推理同样"
              "串行。HTTP 层保持并发。识别结果按图像 SHA256 做 128 条内存 LRU 缓存，命中约 0.01 s。")

    h1(doc, "接口", "03")
    table(doc,
          ["接口", "方法", "说明"],
          [
              ["/api/health", "GET", "健康状态、模型是否加载"],
              ["/api/recognize", "POST", "图片 → 多结构 SMILES（single=true 强制单结构）"],
              ["/api/analyze", "POST", "SMILES → 2D/3D/描述符，n_conformers 可选 1–20"],
              ["/api/pdf", "POST", "PDF → 逐页识别结果"],
              ["/samples/*、/static/*", "GET", "示例图与前端资源"],
          ],
          [Cm(3.6), Cm(1.8), Cm(11.2)])

    h1(doc, "评测方法与结果", "04")
    para(doc, "识别准确率按 InChIKey 集合比较（Markush 等为空时退回 canonical SMILES）；"
              "Golden set 14 个用例，低清为软用例不计入严格分。")
    table(doc,
          ["测试", "规模", "结果", "说明"],
          [
              ["Golden set", "14 用例", "13/13 严格命中", "TP=15，FP=0，FN=0，F1=1.00"],
              ["NCI 渲染压测", "100 分子", "98/100 精确", "connectivity 100/100"],
              ["多结构图", "2×2、1×3", "4/4、3/3", "自动切分 + 逐结构识别"],
              ["合成反应式", "3 结构", "3/3", "箭头不产生伪结构"],
              ["3D 气相精度", "12 个 PDB 配体", "中位 1.82 → 0.90 Å", "单构象 → 20 构象"],
              ["3D 规模", "4,999 分子", "4,978 成功（99.6%）", "21 个失败均为金属有机体系"],
              ["回归", "23 项", "23/23", "健康、缓存、并发、Markush、缩写"],
          ],
          [Cm(3.0), Cm(2.8), Cm(4.2), Cm(6.6)])
    para(doc, "3D 口径说明：上述 RMSD 是气相力场下的构象生成误差，不含蛋白环境，也未做口袋约束优化。",
         size=9, color=MUTED)

    h1(doc, "部署与复现", "05")
    para(doc, "环境：Windows、Python 3.12、RDKit 2026.09.1、DECIMER 2.8 + TensorFlow 2.20；"
              "服务默认监听 127.0.0.1:8765。")
    for line in [
        r"cd molvision",
        r".\run.ps1                       # 启动，首次预热约 3.5 min",
        r"python scripts\selfcheck.py     # 23 项回归",
        r"python scripts\evaluate.py      # Golden 14 用例",
        r"python scripts\pdb_rmsd.py --confs 20",
        r"python scripts\run_pdf.py paper.pdf --out data\pdf_out",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(1)
        r = p.add_run(line)
        run_font(r, 8.5, False, RGBColor(0x33, 0x33, 0x33), font="Consolas")

    h1(doc, "已知边界", "06")
    table(doc,
          ["边界", "表现", "应对"],
          [
              ["真实密集反应式", "自动切分可能拆错", "Crop 手动框选；演示用已验证图"],
              ["手绘 / 低清", "准确率下降", "手绘模型回退 + 人工核对"],
              ["E/Z 过度指定", "结构对、多标构型", "stereo 标记，不自动剥离"],
              ["金属配合物 3D", "UFF 几何不可靠", "仅用于看连接关系"],
              ["Markush R 未定义", "统一按 CH3", "界面标注，可手动改 SMILES"],
          ],
          [Cm(3.6), Cm(5.6), Cm(7.4)])

    h1(doc, "仓库结构", "07")
    for line in [
        "app/        segment.py  recognize.py  chem.py  pdf.py  server.py",
        "web/        index.html（单页 UI，含 Crop / 3D / 导出）",
        "scripts/    evaluate.py  selfcheck.py  pdb_rmsd.py  stress_*.py",
        "samples/    演示图与 14 个 Golden 用例",
        "data/       评测报告、压测日志、截图",
    ]:
        p = doc.add_paragraph()
        p.paragraph_format.left_indent = Cm(0.4)
        p.paragraph_format.space_after = Pt(2)
        r = p.add_run(line)
        run_font(r, 9, False, INK, font="Consolas")

    doc.save(OUT / "02-技术文档.docx")
    print("saved 02")


# ---------------------------------------------------------------------------
# 03 第三方资源说明
# ---------------------------------------------------------------------------
def build_thirdparty() -> None:
    doc = Document()
    setup(doc)
    header_footer(doc, "第三方资源说明")
    doc_title(doc, "第三方资源说明", "依赖、模型、数据与许可证", "TECHATHON 2026  ·  项目内容")

    h1(doc, "使用原则", "01")
    para(doc, "本项目仅通过公开包管理器安装并调用第三方库，未修改其源代码；模型权重通过官方发布渠道"
              "下载；数据来自公开数据集。以下逐项列出组件、版本、许可证与用途。")

    h1(doc, "软件依赖", "02")
    table(doc,
          ["组件", "版本", "许可证", "用途"],
          [
              ["DECIMER", "2.8.0", "MIT", "结构式图像识别（预训练模型）"],
              ["RDKit", "2026.09.1", "BSD-3-Clause", "SMILES 校验、描述符、2D/3D、SDF"],
              ["TensorFlow", "2.20.0", "Apache-2.0", "DECIMER 推理后端"],
              ["Keras", "3.15.1", "Apache-2.0", "DECIMER 模型加载"],
              ["SELFIES", "2.2.0", "Apache-2.0", "分子字符串编码（DECIMER 依赖）"],
              ["OpenCV (headless)", "5.0.0", "Apache-2.0", "图像预处理、切分"],
              ["PyMuPDF", "1.28.2", "AGPL-3.0 或商业许可", "PDF 页面渲染（见合规提示）"],
              ["Pillow", "13.0.0", "HPND", "图像读写与格式转换"],
              ["NumPy", "2.5.4", "BSD-3-Clause", "数值计算"],
              ["pandas", "3.0.6", "BSD-3-Clause", "CSV 汇总"],
              ["3Dmol.js", "本地 vendor", "BSD-3-Clause", "前端 3D 分子可视化"],
          ],
          [Cm(3.2), Cm(2.4), Cm(4.4), Cm(6.6)])

    h1(doc, "模型与数据", "03")
    table(doc,
          ["资源", "来源", "说明"],
          [
              ["DECIMER 权重", "Zenodo / DECIMER 官方发布", "预训练模型，MIT 许可；未参与训练"],
              ["NCI 分子子集", "RDKit 自带数据", "用于识别压测，公开数据"],
              ["PDB 晶体结构", "RCSB PDB", "公开结构，用于 3D RMSD 验证"],
              ["CCD 配体 SMILES", "RCSB Chemical Component Dictionary", "公开数据"],
              ["示例结构图", "团队自行渲染 / 自绘", "不随仓库分发他人受版权图片"],
          ],
          [Cm(3.4), Cm(7.0), Cm(6.2)])

    h1(doc, "合规提示", "04")
    bullet(doc, "PyMuPDF 采用 AGPL-3.0 或商业双许可：本项目以源码形式调用、未修改其代码；"
                "如后续需要闭源分发，应购买商业许可或替换为 pypdf/pdfium 等方案。")
    bullet(doc, "DECIMER 为 MIT 许可；其模型权重的使用请同时遵循 Zenodo 发布页的说明。")
    bullet(doc, "3Dmol.js 为 BSD-3-Clause，前端以本地文件方式引入。")
    bullet(doc, "本项目代码建议采用 MIT 许可证；第三方组件保留各自许可证与版权声明。")

    h1(doc, "原创性边界", "05")
    para(doc, "团队原创部分：切分与后处理、化学归一规则（R 基/缩写/同位素氢）、识别兜底与合并、"
              "3D/多构象流程、Web UI、评测脚本与全部文档。")
    para(doc, "调用部分：DECIMER 识别模型、RDKit 化学计算、TensorFlow 推理、PyMuPDF 渲染、"
              "3Dmol.js 可视化。上述调用不改变第三方组件的许可与版权归属。")

    doc.save(OUT / "03-第三方资源说明.docx")
    print("saved 03")


if __name__ == "__main__":
    build_intro()
    build_tech()
    build_thirdparty()
